import openpyxl

def main():
    file_path = r"D:\ASUS\News & Report\News\News.xlsx"
    print(f"正在開啟 Excel 檔案: {file_path}")
    
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
    except Exception as e:
        print(f"開啟檔案失敗: {e}")
        print("請確認 Excel 檔案沒有被其他程式（如 Excel 軟體）開啟。")
        return

    insights = {
        "PC產業": "[AI觀點] 獨立AI硬體與NPU的發展正重塑終端運算架構，同時中國商用PC顯示器市場展現韌性。華碩可積極佈局獨立NPU產品線與商用顯示方案，以把握AI邊緣運算及企業升級的新商機。",
        "上游廠商動態": "[AI觀點] 記憶體與CPU供應吃緊恐延續至2027年，且次世代GPU時程推遲，將限制PC硬體大幅升級的空間與降價彈性。華碩需提前鞏固關鍵零組件料源，並透過強化軟體整合來提升現有架構產品的附加價值。",
        "各品牌廠動態": "[AI觀點] 五大品牌廠同步推出Googlebook，宣告AI PC進入新一輪軟硬體生態系的激烈戰局；競爭對手更在伺服器市場取得領先。華碩應藉此換機潮凸顯自家Googlebook優勢，並同步強化商用伺服器佈局。",
        "國際經濟": "[AI觀點] 國際油價與美債殖利率飆升使高利率環境延續，進一步壓抑全球終端消費與企業投資動能。面對總體經濟的諸多不確定性，華碩需對第四季銷售目標保持審慎，並強化成本與現金流控管。"
    }

    # 尋找對應的欄位索引 (1-based)
    headers = {cell.value: i for i, cell in enumerate(ws[1], 1)}
    year_col = headers.get('Year')
    week_col = headers.get('Week')
    cat_col = headers.get('Category')
    impact_col = headers.get('對於筆電市場/華碩的影響')

    if not all([year_col, week_col, cat_col, impact_col]):
        print("找不到對應的欄位，請確認 Excel 第一列是否有 'Year', 'Week', 'Category', '對於筆電市場/華碩的影響' 標題。")
        return

    updated_count = 0
    for row in range(2, ws.max_row + 1):
        year = ws.cell(row=row, column=year_col).value
        week = ws.cell(row=row, column=week_col).value
        category = ws.cell(row=row, column=cat_col).value
        
        # 轉換為數字以防格式差異
        try:
            year = int(year) if year is not None else None
            week = int(week) if week is not None else None
        except ValueError:
            pass

        # 改為 2026 年第 39 週
        if year == 2026 and week == 39 and category in insights:
            ws.cell(row=row, column=impact_col).value = insights[category]
            updated_count += 1

    print(f"成功找到並填入了 {updated_count} 筆 wk39 的 AI 觀點。")
    
    try:
        print("正在儲存檔案...")
        wb.save(file_path)
        print("✅ Excel 檔案更新成功！")
    except Exception as e:
        print(f"儲存檔案失敗: {e}")
        print("請確認 Excel 檔案沒有被其他程式開啟，再試一次。")

if __name__ == "__main__":
    main()
